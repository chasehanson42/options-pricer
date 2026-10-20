"""Export Markdown documents to .md, .pdf (reportlab) and .docx (python-docx).

Supported Markdown subset: '#', '##', '###', '- ' / '* ' bullets, paragraphs,
blank lines, **bold** and *italic*. Both renderers share `parse_markdown`.
"""
from __future__ import annotations

import re
from pathlib import Path
from xml.sax.saxutils import escape

SUPPORTED_FORMATS = ("md", "pdf", "docx")

_BULLET_RE = re.compile(r"^\s*[-*]\s+(.*)$")
_HEADING_RE = re.compile(r"^(#{1,3})\s+(.*?)\s*#*\s*$")
# Inline tokens: **bold** or *italic*.
_INLINE_RE = re.compile(r"\*\*(.+?)\*\*|(?<![*\w])\*(?!\s)(.+?)(?<!\s)\*(?![*\w])")


def parse_markdown(md: str) -> list[tuple[str, str]]:
    """Parse the supported Markdown subset into (kind, text) blocks.

    kinds: h1, h2, h3, bullet, para. Consecutive non-blank text lines are
    joined into one paragraph with explicit line breaks preserved as '\\n'.
    """
    blocks: list[tuple[str, str]] = []
    para: list[str] = []

    def flush() -> None:
        if para:
            blocks.append(("para", "\n".join(para)))
            para.clear()

    for raw in (md or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush()
            continue
        if line.strip().startswith("```"):
            continue  # never render stray fences
        hm = _HEADING_RE.match(line.strip())
        if hm:
            flush()
            blocks.append((f"h{len(hm.group(1))}", hm.group(2)))
            continue
        bm = _BULLET_RE.match(line)
        if bm and not line.strip().startswith("**"):
            flush()
            blocks.append(("bullet", bm.group(1).strip()))
            continue
        para.append(line.strip())
    flush()
    return blocks


def inline_segments(text: str) -> list[tuple[str, bool, bool]]:
    """Split text into (segment, bold, italic) runs."""
    out: list[tuple[str, bool, bool]] = []
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False, False))
        if m.group(1) is not None:
            out.append((m.group(1), True, False))
        else:
            out.append((m.group(2), False, True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False, False))
    return out


def _to_reportlab_markup(text: str) -> str:
    parts = []
    for seg, bold, italic in inline_segments(text):
        s = escape(seg).replace("\n", "<br/>")
        if bold:
            s = f"<b>{s}</b>"
        if italic:
            s = f"<i>{s}</i>"
        parts.append(s)
    return "".join(parts)


def _write_pdf(blocks: list[tuple[str, str]], path: Path) -> None:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    base = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=base["Normal"], fontName="Helvetica",
                          fontSize=10.5, leading=14, alignment=TA_LEFT, spaceAfter=6)
    styles = {
        "h1": ParagraphStyle("H1", parent=body, fontName="Helvetica-Bold", fontSize=18,
                             leading=22, spaceAfter=4),
        "h2": ParagraphStyle("H2", parent=body, fontName="Helvetica-Bold", fontSize=13,
                             leading=16, spaceBefore=10, spaceAfter=4),
        "h3": ParagraphStyle("H3", parent=body, fontName="Helvetica-Bold", fontSize=11,
                             leading=14, spaceBefore=6, spaceAfter=2),
        "para": body,
        "bullet": ParagraphStyle("Bullet", parent=body, spaceAfter=2),
    }

    story: list = []
    pending_bullets: list = []

    def flush_bullets() -> None:
        if pending_bullets:
            story.append(ListFlowable(
                [ListItem(p, leftIndent=12) for p in pending_bullets],
                bulletType="bullet", start="•", leftIndent=12, bulletFontName="Helvetica"))
            story.append(Spacer(1, 4))
            pending_bullets.clear()

    for kind, text in blocks:
        para = Paragraph(_to_reportlab_markup(text), styles[kind])
        if kind == "bullet":
            pending_bullets.append(para)
            continue
        flush_bullets()
        story.append(para)
    flush_bullets()
    if not story:
        story.append(Spacer(1, 1))

    margin = 0.75 * inch
    doc = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=margin)
    doc.build(story)


def _add_runs(paragraph, text: str) -> None:
    for seg, bold, italic in inline_segments(text):
        lines = seg.split("\n")
        for i, chunk in enumerate(lines):
            run = paragraph.add_run(chunk)
            run.bold = bold or None
            run.italic = italic or None
            if i < len(lines) - 1:
                run.add_break()


def _write_docx(blocks: list[tuple[str, str]], path: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    for section in document.sections:
        section.left_margin = section.right_margin = Inches(0.75)
        section.top_margin = section.bottom_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    levels = {"h1": 1, "h2": 2, "h3": 3}
    for kind, text in blocks:
        if kind in levels:
            p = document.add_heading("", level=levels[kind])
        elif kind == "bullet":
            p = document.add_paragraph(style="List Bullet")
        else:
            p = document.add_paragraph()
        _add_runs(p, text)
    document.save(str(path))


def export_documents(markdown: str, out_stem: Path, formats: list[str]) -> list[Path]:
    """Write `markdown` as each requested format next to `out_stem`; return paths."""
    out_stem = Path(out_stem)
    wanted = [f.lower().lstrip(".") for f in formats]
    unknown = [f for f in wanted if f not in SUPPORTED_FORMATS]
    if unknown:
        raise ValueError(f"Unsupported export format(s): {', '.join(unknown)}; "
                         f"supported: {', '.join(SUPPORTED_FORMATS)}")
    out_stem.parent.mkdir(parents=True, exist_ok=True)
    blocks = parse_markdown(markdown)
    written: list[Path] = []
    for fmt in dict.fromkeys(wanted):
        path = out_stem.with_name(out_stem.name + f".{fmt}")
        if fmt == "md":
            path.write_text(markdown if markdown.endswith("\n") else markdown + "\n",
                            encoding="utf-8")
        elif fmt == "pdf":
            _write_pdf(blocks, path)
        else:
            _write_docx(blocks, path)
        written.append(path)
    return written

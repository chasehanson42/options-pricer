"""Tests for the Markdown -> md/pdf/docx exporter."""
from __future__ import annotations

import pytest

from exporters.pdf_builder import export_documents, inline_segments, parse_markdown

SAMPLE = """# Jane Doe
Austin, TX | jane.doe@example.com

## Summary
Backend engineer with **6 years** of *Python* & <SQL> work.

## Experience
### Senior Backend Engineer — Acme Analytics
Remote | 2022-03 – present
- Designed a **FastAPI** service.
* Cut latency from 900ms to 180ms.
"""


def test_parse_markdown_blocks():
    blocks = parse_markdown(SAMPLE)
    kinds = [k for k, _ in blocks]
    assert kinds == ["h1", "para", "h2", "para", "h2", "h3", "para", "bullet", "bullet"]
    assert blocks[0] == ("h1", "Jane Doe")
    assert blocks[7] == ("bullet", "Designed a **FastAPI** service.")


def test_inline_segments():
    segs = inline_segments("a **b** *c* d")
    assert segs == [("a ", False, False), ("b", True, False), (" ", False, False),
                    ("c", False, True), (" d", False, False)]


def test_export_all_formats(tmp_path):
    paths = export_documents(SAMPLE, tmp_path / "out" / "resume", ["md", "pdf", "docx"])
    assert [p.suffix for p in paths] == [".md", ".pdf", ".docx"]
    for p in paths:
        assert p.exists() and p.stat().st_size > 0
    assert paths[0].read_text(encoding="utf-8").startswith("# Jane Doe")
    assert paths[1].read_bytes().startswith(b"%PDF")

    import docx
    d = docx.Document(str(paths[2]))
    styles = [(p.style.name, p.text) for p in d.paragraphs]
    assert ("Heading 1", "Jane Doe") in styles
    assert ("Heading 3", "Senior Backend Engineer — Acme Analytics") in styles
    assert any(s == "List Bullet" and "FastAPI" in t for s, t in styles)
    bullet = next(p for p in d.paragraphs if "FastAPI" in p.text)
    assert any(r.bold and r.text == "FastAPI" for r in bullet.runs)


def test_unknown_format_raises(tmp_path):
    with pytest.raises(ValueError):
        export_documents("# x", tmp_path / "x", ["md", "html"])
